"""
tools/file_handler.py
Xử lý file Excel, PDF và tạo QR code cho tài sản
"""
import io
import logging
import qrcode
from PIL import Image
from datetime import date

logger = logging.getLogger(__name__)


# ── Excel Import ──────────────────────────────────────────

def parse_assets_from_excel(file_bytes: bytes) -> list:
    """
    Đọc file Excel danh sách tài sản.
    Trả về list dict, hoặc [] nếu lỗi.
    Cột mong đợi: Tên tài sản, Loại, Serial, Ngày mua, Hết bảo hành, Trạng thái, Phòng ban, Người dùng, Vị trí, Giá mua, Ghi chú
    """
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)
        ws = wb.active
        headers = []
        assets = []
        col_map = {
            "tên tài sản": "name", "tên": "name",
            "loại": "category", "loại tài sản": "category",
            "serial": "serial_number", "số serial": "serial_number",
            "ngày mua": "purchase_date",
            "hết bảo hành": "warranty_end_date", "ngày hết bảo hành": "warranty_end_date",
            "trạng thái": "status",
            "phòng ban": "department",
            "người dùng": "assigned_to", "nhân viên": "assigned_to",
            "vị trí": "location",
            "giá mua": "purchase_price", "giá": "purchase_price",
            "ghi chú": "notes",
        }

        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i == 0:
                headers = [str(c).strip().lower() if c else "" for c in row]
                continue
            if not any(row):
                continue
            asset = {}
            for j, val in enumerate(row):
                if j >= len(headers):
                    break
                h = headers[j]
                field = col_map.get(h)
                if field and val is not None:
                    # Xử lý ngày tháng
                    if field in ("purchase_date", "warranty_end_date"):
                        if hasattr(val, "date"):
                            val = val.date().isoformat()
                        elif isinstance(val, str):
                            val = val.strip()
                    elif field == "purchase_price":
                        try:
                            val = float(str(val).replace(",", "").replace(".", "").strip() or 0)
                        except Exception:
                            val = 0
                    else:
                        val = str(val).strip()
                    asset[field] = val
            if asset.get("name"):
                assets.append(asset)

        logger.info(f"Đọc được {len(assets)} tài sản từ Excel.")
        return assets
    except Exception as e:
        logger.error(f"Lỗi parse_assets_from_excel: {e}")
        return []


def export_assets_to_excel(assets: list) -> bytes:
    """Xuất danh sách tài sản ra file Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Danh sách tài sản"

        headers = [
            ("ID", 10), ("Tên tài sản", 30), ("Loại", 15),
            ("Serial", 18), ("Ngày mua", 14), ("Hết bảo hành", 14),
            ("Trạng thái", 18), ("Phòng ban", 18), ("Người dùng", 20),
            ("Vị trí", 15), ("Giá mua (VNĐ)", 18), ("Ghi chú", 25),
        ]
        fields = [
            "id", "name", "category", "serial_number", "purchase_date",
            "warranty_end_date", "status", "department", "assigned_to",
            "location", "purchase_price", "notes"
        ]

        header_fill = PatternFill("solid", fgColor="1E3A5F")
        header_font = Font(bold=True, color="FFFFFF", size=11)
        thin = Side(style="thin", color="CCCCCC")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)

        for col, (h, w) in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border
            ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = w

        alt_fill = PatternFill("solid", fgColor="F0F4F8")
        for row_idx, asset in enumerate(assets, 2):
            fill = alt_fill if row_idx % 2 == 0 else None
            for col, field in enumerate(fields, 1):
                val = asset.get(field, "")
                if val is None:
                    val = ""
                cell = ws.cell(row=row_idx, column=col, value=val)
                cell.border = border
                if fill:
                    cell.fill = fill

        ws.row_dimensions[1].height = 24
        ws.freeze_panes = "A2"

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as e:
        logger.error(f"Lỗi export_assets_to_excel: {e}")
        return b""


def export_warranty_report(assets: list) -> bytes:
    """Xuất báo cáo bảo hành ra Excel."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Báo cáo bảo hành"
        ws["A1"] = "BÁO CÁO TÀI SẢN SẮP HẾT BẢO HÀNH"
        ws["A1"].font = Font(bold=True, size=14, color="CC0000")
        ws["A2"] = f"Ngày xuất: {date.today().strftime('%d/%m/%Y')}"

        headers = ["ID", "Tên tài sản", "Loại", "Người dùng", "Phòng ban", "Hết bảo hành", "Trạng thái"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="CC4400")

        fields = ["id", "name", "category", "assigned_to", "department", "warranty_end_date", "status"]
        for ri, a in enumerate(assets, 5):
            for ci, f in enumerate(fields, 1):
                ws.cell(row=ri, column=ci, value=a.get(f, ""))

        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    except Exception as e:
        logger.error(f"Lỗi export_warranty_report: {e}")
        return b""


# ── QR Code ───────────────────────────────────────────────

def generate_qr_code(asset_id: str, asset_name: str) -> bytes:
    """
    Tạo QR code cho tài sản.
    QR chứa thông tin ID + tên, trả về bytes PNG.
    """
    try:
        content = f"ASSET:{asset_id}|{asset_name}"
        qr = qrcode.QRCode(
            version=2,
            error_correction=qrcode.constants.ERROR_CORRECT_M,
            box_size=8,
            border=3,
        )
        qr.add_data(content)
        qr.make(fit=True)
        img = qr.make_image(fill_color="#1E3A5F", back_color="white")

        # Thêm label dưới QR
        from PIL import ImageDraw, ImageFont
        img_pil = img.convert("RGB")
        w, h = img_pil.size
        label_h = 40
        new_img = Image.new("RGB", (w, h + label_h), "white")
        new_img.paste(img_pil, (0, 0))
        draw = ImageDraw.Draw(new_img)
        text = f"{asset_id}"
        draw.text((w // 2, h + label_h // 2), text, fill="#1E3A5F", anchor="mm")

        buf = io.BytesIO()
        new_img.save(buf, format="PNG")
        return buf.getvalue()
    except Exception as e:
        logger.error(f"Lỗi generate_qr_code: {e}")
        # Fallback: QR đơn giản
        try:
            qr = qrcode.make(f"ASSET:{asset_id}")
            buf = io.BytesIO()
            qr.save(buf, format="PNG")
            return buf.getvalue()
        except Exception:
            return b""


def generate_asset_label_pdf(asset: dict) -> bytes:
    """Tạo nhãn PDF cho tài sản (dán trên thiết bị)."""
    try:
        from reportlab.lib.pagesizes import A6, landscape
        from reportlab.lib.units import cm
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Image as RLImage, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
        from reportlab.lib import colors

        buf = io.BytesIO()
        doc = SimpleDocTemplate(buf, pagesize=landscape(A6), topMargin=0.5*cm, bottomMargin=0.5*cm,
                                leftMargin=0.5*cm, rightMargin=0.5*cm)

        styles = getSampleStyleSheet()
        title_style = ParagraphStyle("title", fontSize=12, fontName="Helvetica-Bold",
                                     textColor=colors.HexColor("#1E3A5F"))
        small_style = ParagraphStyle("small", fontSize=8, fontName="Helvetica")

        qr_bytes = generate_qr_code(asset.get("id", ""), asset.get("name", ""))
        qr_img = RLImage(io.BytesIO(qr_bytes), width=3*cm, height=3*cm)

        info = [
            Paragraph(f"<b>{asset.get('name', '')}</b>", title_style),
            Paragraph(f"ID: {asset.get('id', '')}", small_style),
            Paragraph(f"Loại: {asset.get('category', '')}", small_style),
            Paragraph(f"Serial: {asset.get('serial_number', '')}", small_style),
            Paragraph(f"Phòng: {asset.get('department', '')}", small_style),
        ]

        table = Table([[qr_img, info]], colWidths=[3.5*cm, 10*cm])
        table.setStyle(TableStyle([
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ]))

        doc.build([table])
        return buf.getvalue()
    except Exception as e:
        logger.error(f"Lỗi generate_asset_label_pdf: {e}")
        return b""
